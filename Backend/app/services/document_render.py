"""Render generated Markdown documents to Word (.docx).

Deliberately GENERIC. Minutes of Meeting is the first caller, but the requirement document and
the BRD/FRD will use the same function — only the title, subtitle and body differ. Do not add
MOM-specific logic here; pass it in.

Scope of Markdown supported. Measured against the real stored minutes rather than assumed:

    headings   # ## ###          (levels 1-6 accepted)
    bullets    - item            (nesting by indent)
    rules      ---
    inline     `code`  **bold**  *italic*
    numbered   1. item
    paragraphs anything else

The live MOM prompt currently emits only headings, bullets, rules and inline code — no bold,
tables or numbered lists. The extra cases are handled anyway because prompts change, and a
prompt edit should never silently produce a broken document.

NOT supported: tables, images, links, blockquotes, fenced code blocks. If a prompt starts
emitting those, they fall through to plain paragraphs rather than breaking the render.
"""
from __future__ import annotations

import io
import logging
import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

logger = logging.getLogger(__name__)

# Matches the blue used by the web UI, so a downloaded document and the screen look related.
_ACCENT = RGBColor(0x0D, 0x76, 0xFF)
_MUTED = RGBColor(0x66, 0x66, 0x66)
_RULE_COLOR = "CCCCCC"

# Company mark, placed at the top of every generated document. Lives under app/assets so it
# ships with the backend; the repo-root copy is not deployed.
_LOGO_PATH = Path(__file__).parent.parent / "assets" / "nfclogo.jpg"
_LOGO_WIDTH = Inches(0.8)   # square source (200x200) -> ~250 DPI on the page, sharp in print

# One typeface across the whole document. Without this, python-docx's bundled template applies
# the Office 2007 theme, which sets headings in Calibri and BODY TEXT IN CAMBRIA — the reverse
# of what people expect, and visibly inconsistent next to a branded header.
_FONT = "Calibri"
_STYLED = (
    "Normal", "Title", "Heading 1", "Heading 2", "Heading 3",
    "Heading 4", "Heading 5", "Heading 6",
    "List Bullet", "List Bullet 2", "List Bullet 3", "List Number",
)


def _apply_font(doc: Document) -> None:
    """Force one typeface across every style the renderer uses."""
    for style_name in _STYLED:
        try:
            style = doc.styles[style_name]
        except KeyError:
            continue                                  # style absent in this template
        style.font.name = _FONT
        # .font.name sets only the Latin run property. Word falls back to the THEME font for
        # east-Asian and complex-script characters unless these are set too, which would leave
        # stray Cambria on any non-Latin character the transcript contains.
        r_pr = style.element.get_or_add_rPr()
        r_fonts = r_pr.find(qn("w:rFonts"))
        if r_fonts is None:
            r_fonts = OxmlElement("w:rFonts")
            r_pr.append(r_fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            r_fonts.set(qn(attr), _FONT)


def _add_logo(doc: Document) -> None:
    """Company mark at the top. Missing file is logged, never fatal — a download must still work."""
    if not _LOGO_PATH.exists():
        logger.warning(f"Logo not found at {_LOGO_PATH}; generating document without it.")
        return
    try:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.add_run().add_picture(str(_LOGO_PATH), width=_LOGO_WIDTH)
    except Exception as exc:
        logger.warning(f"Could not embed the logo ({exc}); continuing without it.")

_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET_RE = re.compile(r"^([-*+])\s+(.*)$")
_NUMBERED_RE = re.compile(r"^\d+[.)]\s+(.*)$")
_RULE_RE = re.compile(r"^([-*_])\1{2,}$")
# Split on inline spans while KEEPING the delimiters, so runs can be styled individually.
_INLINE_RE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|(?<!\*)\*[^*]+\*(?!\*))")


def _add_bottom_border(paragraph) -> None:
    """A horizontal rule. python-docx has no native one, so use a paragraph bottom border."""
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), _RULE_COLOR)
    borders.append(bottom)
    p_pr.append(borders)


def _add_inline(paragraph, text: str) -> None:
    """Add text to a paragraph, styling `code`, **bold** and *italic* spans as separate runs."""
    for part in _INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _bullet_style(indent_spaces: int) -> str:
    """Word ships List Bullet, List Bullet 2 and List Bullet 3. Map indent depth onto them."""
    level = min(indent_spaces // 2, 2)
    return "List Bullet" if level == 0 else f"List Bullet {level + 1}"


def render_markdown_to_docx(
    markdown: str,
    title: str,
    subtitle: str = "",
    footer: str = "",
    warning: str = "",
) -> bytes:
    """
    Render Markdown to a Word document and return it as bytes.

    title    large heading at the top, e.g. "Minutes of Meeting"
    subtitle grey line beneath it, e.g. "05-04-26_GROOMING · 4 May 2026"
    warning  optional highlighted banner, e.g. an incomplete-document notice
    footer   small grey line at the end

    Returns bytes so the caller can stream the file without writing to disk.
    Raises nothing for empty content — produces a document saying so, because a caller
    downloading a file should get a file, not a stack trace.
    """
    doc = Document()
    _apply_font(doc)

    # ── Header block ────────────────────────────────────────────────────────
    _add_logo(doc)

    heading = doc.add_heading(title, level=0)
    for run in heading.runs:
        run.font.color.rgb = _ACCENT

    if subtitle:
        para = doc.add_paragraph()
        run = para.add_run(subtitle)
        run.font.size = Pt(10)
        run.font.color.rgb = _MUTED
        _add_bottom_border(para)

    if warning:
        para = doc.add_paragraph()
        run = para.add_run(warning)
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x96, 0x59, 0x0C)  # amber, matches the UI warning

    if not (markdown or "").strip():
        doc.add_paragraph("This document has no content.")
        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    # ── Body ────────────────────────────────────────────────────────────────
    for raw_line in markdown.split("\n"):
        line = raw_line.rstrip()
        stripped = line.strip()

        if not stripped:
            continue  # Word spaces paragraphs itself; blank lines would double it

        if _RULE_RE.match(stripped):
            _add_bottom_border(doc.add_paragraph())
            continue

        heading_match = _HEADING_RE.match(stripped)
        if heading_match:
            level = min(len(heading_match.group(1)), 6)
            _add_inline(doc.add_heading(level=level), heading_match.group(2))
            continue

        bullet_match = _BULLET_RE.match(stripped)
        if bullet_match:
            indent = len(line) - len(line.lstrip())
            _add_inline(
                doc.add_paragraph(style=_bullet_style(indent)),
                bullet_match.group(2),
            )
            continue

        numbered_match = _NUMBERED_RE.match(stripped)
        if numbered_match:
            _add_inline(doc.add_paragraph(style="List Number"), numbered_match.group(1))
            continue

        _add_inline(doc.add_paragraph(), stripped)

    # ── Footer block ────────────────────────────────────────────────────────
    if footer:
        para = doc.add_paragraph()
        _add_bottom_border(para)
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(footer)
        run.font.size = Pt(8)
        run.font.color.rgb = _MUTED

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def safe_filename(*parts: str, extension: str = "docx") -> str:
    """
    Build a download filename that survives Windows, macOS and Linux.

    safe_filename("MOM", "05-04-26_GROOMING") -> "MOM_05-04-26_GROOMING.docx"
    """
    cleaned = []
    for part in parts:
        text = re.sub(r"[^A-Za-z0-9._-]+", "_", (part or "").strip())
        text = text.strip("_")
        if text:
            cleaned.append(text)
    stem = "_".join(cleaned) or f"document_{datetime.now():%Y%m%d}"
    return f"{stem[:120]}.{extension}"
